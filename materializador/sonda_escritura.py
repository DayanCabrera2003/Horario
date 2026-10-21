"""Escribe en el libro como lo haria quien lo usa, recalcula y comprueba.

EL PROBLEMA QUE RESUELVE

Los verificadores comparan el libro tal como sale del generador. Pero el
libro es un instrumento: alguien va a escribir en el. Dos fallos que encontro
la auditoria del 2026-10-17 solo aparecian al escribir:

  - una fecha tecleada de verdad (un numero de serie, no un texto ISO)
    comparaba mal con una fecha escrita en la descripcion;
  - en una coleccion que crece, una fila nueva no calculaba nada, porque sus
    campos fijos estaban bloqueados y la guarda solo miraba el primero.

Esta sonda escribe valores en celdas, recalcula con LibreOffice y comprueba
el valor de otras celdas. La usan ./demostrar.sh y ./demostrar-robustez.sh.

Uso:
    python3 sonda_escritura.py <libro.xlsx> <hoja> <escritura>... -- <comprobacion>...

Cada escritura es CELDA=VALOR; un VALOR "fecha:AAAA-MM-DD" se escribe como
fecha de verdad, uno numerico como numero y lo demas como texto. Cada
comprobacion es CELDA=ESPERADO.

Una CELDA puede llevar delante otra hoja, como Parametros!B4: un parametro
se escribe en su hoja y lo que cambia se lee en otra. Sin hoja, es la
<hoja> de la linea de ordenes.
"""

import datetime
import shutil
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from verificar import recalcular_con_libreoffice


def valor_de(texto):
    """El valor que se escribe a partir de su forma en la linea de ordenes."""
    if texto.startswith("fecha:"):
        return datetime.date.fromisoformat(texto[len("fecha:"):])
    try:
        return int(texto)
    except ValueError:
        try:
            return float(texto)
        except ValueError:
            return texto


def partir_celda(texto, hoja):
    """(HOJA, CELDA) de una celda escrita como CELDA o como HOJA!CELDA."""
    if "!" in texto:
        otra, celda = texto.rsplit("!", 1)
        return otra, celda
    return hoja, texto


def main(argv):
    if len(argv) < 5 or "--" not in argv:
        print(__doc__)
        return 2
    libro, hoja = argv[1], argv[2]
    corte = argv.index("--")
    escrituras = [a.split("=", 1) for a in argv[3:corte]]
    comprobaciones = [a.split("=", 1) for a in argv[corte + 1:]]
    temporal = Path(tempfile.mkdtemp(prefix="situacion-sonda-"))
    try:
        copia = temporal / Path(libro).name
        wb = load_workbook(libro)
        for celda, valor in escrituras:
            donde, celda = partir_celda(celda, hoja)
            wb[donde][celda].value = valor_de(valor)
        wb.save(copia)
        recalculado, otro = recalcular_con_libreoffice(copia)
        try:
            leido = load_workbook(recalculado, data_only=True)
            obtenidos = {}
            for celda, _ in comprobaciones:
                donde, sola = partir_celda(celda, hoja)
                obtenidos[celda] = leido[donde][sola].value
        finally:
            shutil.rmtree(otro, ignore_errors=True)
    finally:
        shutil.rmtree(temporal, ignore_errors=True)
    descripcion = "{}: escribir {}".format(
        hoja, ", ".join("{}={}".format(c, v) for c, v in escrituras))
    fallos = ["{} vale {!r} y no {!r}".format(celda, obtenidos[celda], esperado)
              for celda, esperado in comprobaciones
              if str(obtenidos[celda]) != esperado]
    if fallos:
        print("FALLA  {}: {}".format(descripcion, "; ".join(fallos)))
        return 1
    print("  ok   {}: {}".format(descripcion, ", ".join(
        "{}={}".format(c, obtenidos[c]) for c, _ in comprobaciones)))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
