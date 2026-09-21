"""Convierte un plano de libro en un archivo .xlsx.

Este programa no sabe nada del dominio. No conoce profesores, ni defensas, ni
asignaturas: conoce celdas, formulas, estilos y rangos. Todo el trabajo
semantico ya lo hizo la arquitectura de Excel, en Lisp.

POR QUE EXISTE ESTE ARCHIVO

Porque no hay ninguna libreria de Common Lisp capaz de escribir xlsx con
formulas, formato condicional, validacion de datos, rangos nombrados y
proteccion de celdas. Escribir un emisor OOXML desde cero son semanas de
trabajo que no aportan nada a la pregunta de la tesis, y reintroducirian bugs
que el corpus ya resolvio contra el usuario real.

POR QUE ESTO NO ES EL ERROR DE LA TESIS DE 2026

Alli el JSON con formulas de Excel era el PUNTO DE EXTENSION: el sitio por
donde, segun aquel trabajo, iba a entrar un backend HTML. Por eso la
afirmacion era falsa: lo que ese JSON lleva no es una descripcion del calculo,
es el calculo ya compilado a Excel.

Aqui el punto de extension es el protocolo, que esta cuatro fases mas arriba.
El plano vive dentro de una arquitectura concreta, por debajo de ese punto, y
ninguna otra arquitectura lo toca: la web y el texto plano emiten
directamente, sin intermediario. La asimetria la causa que un .xlsx sea un
contenedor binario comprimido, que es un hecho mecanico y no semantico.

Uso:
    python3 materializar.py <plano.json> <salida.xlsx>
"""

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

# El plano trae las celdas ya bloqueadas por defecto; solo hay que abrir las
# que el rol de campo declaro editables. En OOXML toda celda nace bloqueada y
# ese atributo no hace nada hasta que la hoja se protege.
DESBLOQUEADA = Protection(locked=False)

FUENTE_ENCABEZADO = Font(bold=True, color="FFFFFF")
RELLENO_ENCABEZADO = PatternFill("solid", fgColor="44546A")
LADO_FINO = Side(style="thin", color="BFBFBF")
BORDE = Border(left=LADO_FINO, right=LADO_FINO, top=LADO_FINO, bottom=LADO_FINO)


def materializar(plano, destino):
    """Construye el libro descrito por el plano y lo guarda en destino."""
    libro = Workbook()
    libro.remove(libro.active)

    for hoja_plan in plano["hojas"]:
        hoja = libro.create_sheet(_nombre_valido(hoja_plan["nombre"]))
        _escribir_encabezados(hoja, hoja_plan)
        _escribir_celdas(hoja, hoja_plan)
        _escribir_formulas(hoja, hoja_plan)
        _aplicar_formatos_condicionales(hoja, hoja_plan)
        _aplicar_validaciones(hoja, hoja_plan)
        _escribir_leyenda(hoja, hoja_plan)
        _aplicar_presentacion(hoja, hoja_plan)
        _proteger(hoja, hoja_plan)
        _registrar_rangos(libro, hoja, hoja_plan)

    libro.save(destino)
    return destino


def _nombre_valido(nombre):
    """Los nombres de hoja tienen tope de 31 caracteres y prohiben algunos
    signos. Es una restriccion del formato, no del dominio, y por eso se
    resuelve aqui."""
    limpio = "".join(c for c in nombre if c not in "[]:*?/\\")
    return limpio[:31] or "Hoja"


def _escribir_encabezados(hoja, plan):
    for item in plan.get("encabezados") or []:
        celda = hoja[item["celda"]]
        celda.value = item["texto"]
        celda.font = FUENTE_ENCABEZADO
        celda.fill = RELLENO_ENCABEZADO
        celda.alignment = Alignment(horizontal="center", vertical="center")
    titulo = hoja["A1"]
    titulo.value = plan["nombre"]
    titulo.font = Font(bold=True, size=13)


def _escribir_celdas(hoja, plan):
    for item in plan.get("celdas") or []:
        hoja[item["celda"]] = item["valor"]


def _escribir_formulas(hoja, plan):
    """Las formulas van tambien en las filas de reserva.

    Si no, la primera fila que escriba el usuario no calcularia nada. Es una
    de las cosas que el corpus aprendio por las malas.
    """
    for item in plan.get("formulas") or []:
        hoja[item["celda"]] = item["formula"]


def _aplicar_formatos_condicionales(hoja, plan):
    """Una regla por marca. El color viene del plano, que lo saco de la
    severidad; la marca no conoce ningun color."""
    for item in plan.get("formatos_condicionales") or []:
        relleno = PatternFill("solid", fgColor=item["relleno"])
        regla = FormulaRule(formula=[item["formula"]], fill=relleno, stopIfTrue=False)
        for rango in item["rango"].split():
            hoja.conditional_formatting.add(rango, regla)


def _aplicar_validaciones(hoja, plan):
    """El desplegable de cada campo con dominio.

    `bloquea` sale de :AL-VIOLAR. Cuando es falso se usa un aviso informativo,
    que avisa y deja seguir: es lo que hacen los libros del corpus, y es una
    decision del dominio, no de Excel.
    """
    for item in plan.get("validaciones") or []:
        valores = item.get("valores") or []
        if not valores:
            continue
        formula = '"{}"'.format(",".join(str(v).replace('"', "") for v in valores))
        dv = DataValidation(
            type="list",
            formula1=formula,
            allow_blank=True,
            showErrorMessage=True,
            errorStyle="stop" if item.get("bloquea") else "information",
            error="Ese valor no esta en la lista.",
            errorTitle="Valor fuera del dominio",
        )
        hoja.add_data_validation(dv)
        dv.add(item["rango"])


def _escribir_leyenda(hoja, plan):
    """Que significa cada color, escrito. Un color sin explicacion es medio
    concepto."""
    leyenda = plan.get("leyenda") or []
    if not leyenda:
        return
    fila = plan["fila_primera"] + plan["capacidad"] + 1
    hoja.cell(row=fila, column=1, value="Leyenda").font = Font(bold=True)
    for i, item in enumerate(leyenda, start=1):
        muestra = hoja.cell(row=fila + i, column=1)
        muestra.fill = PatternFill("solid", fgColor=item["relleno"])
        muestra.border = BORDE
        hoja.cell(row=fila + i, column=2, value=item["texto"])


def _aplicar_presentacion(hoja, plan):
    primera, ultima = plan["fila_primera"], plan["fila_primera"] + plan["capacidad"] - 1
    encabezado = plan["fila_encabezado"]
    n_columnas = len(plan.get("encabezados") or []) or 1

    for fila in hoja.iter_rows(min_row=encabezado, max_row=ultima,
                               min_col=1, max_col=n_columnas):
        for celda in fila:
            celda.border = BORDE

    for col in range(1, n_columnas + 1):
        letra = get_column_letter(col)
        ancho = max(
            [len(str(hoja.cell(row=r, column=col).value or ""))
             for r in range(encabezado, min(ultima, primera + 30) + 1)]
            + [10]
        )
        hoja.column_dimensions[letra].width = min(ancho + 3, 42)

    # Inmoviliza el encabezado para que no se pierda al bajar.
    hoja.freeze_panes = hoja.cell(row=primera, column=1)
    if plan.get("color_pestana"):
        hoja.sheet_properties.tabColor = plan["color_pestana"]


def _proteger(hoja, plan):
    """Bloquea todo menos lo que el rol de campo declaro editable.

    Sin contrasena a proposito: el objetivo es evitar que se borre una formula
    sin querer, no proteger secretos.
    """
    editables = plan.get("editables") or []
    for rango in editables:
        minc, minr, maxc, maxr = range_boundaries(rango)
        for fila in hoja.iter_rows(min_row=minr, max_row=maxr,
                                   min_col=minc, max_col=maxc):
            for celda in fila:
                celda.protection = DESBLOQUEADA
    hoja.protection.sheet = True
    hoja.protection.autoFilter = False


def _registrar_rangos(libro, hoja, plan):
    for item in plan.get("rangos_nombrados") or []:
        referencia = item["referencia"].replace(
            "{}!".format(plan["nombre"]), "{}!".format(hoja.title)
        )
        libro.defined_names.add(DefinedName(item["nombre"], attr_text=referencia))


def main(argv):
    if len(argv) != 3:
        print(__doc__)
        return 2
    plano = json.loads(Path(argv[1]).read_text(encoding="utf-8"))
    destino = materializar(plano, argv[2])
    print("Escrito {}".format(destino))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
