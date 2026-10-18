"""Comprueba que las formulas del libro generado EVALUAN, y que dan lo mismo
que el evaluador de referencia.

EL PROBLEMA QUE RESUELVE

Una bateria de pruebas de Python no evalua formulas: comprueba que en la celda
hay la cadena que esperabamos, no que esa cadena calcule bien. Un libro puede
pasar todas las pruebas y abrirse lleno de errores.

EL PROCEDIMIENTO

  1. Se genera el libro.
  2. Se pasa por LibreOffice en modo headless, que lo abre, recalcula y lo
     vuelve a escribir.
  3. Se relee con data_only=True, que da los VALORES y no las formulas.
  4. Se comparan con lo que dijo el evaluador de referencia del nucleo.

Si los dos coinciden, dos cosas quedan demostradas a la vez: que las formulas
son validas, y que la arquitectura de Excel coincide con la semantica del
lenguaje. Ninguna de las tres tesis anteriores de la linea presenta evidencia
de que las formulas que genera se evaluen; la de 2026 lo afirma y lo respalda
con capturas de pantalla, que es lo mas lejos que se llego.

CADA HOJA DICE QUE ES

El plano marca cada hoja con su identidad: "coleccion" en las hojas de datos,
"cruce" en las tablas cruzadas, "relacion" en las auxiliares de una relacion
uno a muchos y "auxiliar" en las hojas de apoyo que no muestran datos. Se empareja por esa
identidad, nunca por los nombres de las columnas: emparejar por nombres hizo
que la hoja principal del horario del Saul -la que lleva columnas de clave
anadidas- se saltara en silencio hasta el 2026-10-17.

CUANDO FALLA

  - Una celda no coincide con el oraculo.
  - Una coleccion del oraculo no tiene hoja, o una hoja no tiene identidad.
  - Una hoja con filas en el oraculo no compara ninguna celda.
  - Cualquier celda del libro, este o no en el oraculo, vale un error de
    formula (#VALUE!, #NAME?, #REF!, #N/A, #DIV/0!...).

Las tablas cruzadas se comprueban contra el oraculo de cruces
(verificar_cruces.py) y las relaciones contra el de valores
(verificar_relaciones.py).

Uso:
    python3 verificar.py <plano.json> <esperado.json> <libro.xlsx> [<cruces.json>]
"""

import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from comparar import iguales, normalizar
from verificar_cruces import verificar_cruces
from verificar_errores import errores_del_libro
from verificar_relaciones import verificar_relacion


def recalcular_con_libreoffice(xlsx):
    """Abre el libro con LibreOffice y lo vuelve a escribir, recalculado."""
    destino = Path(tempfile.mkdtemp(prefix="situacion-"))
    subprocess.run(
        ["soffice", "--headless", "--norestore",
         "--convert-to", "xlsx", "--outdir", str(destino), str(xlsx)],
        check=True, capture_output=True, timeout=180,
    )
    salida = destino / Path(xlsx).name
    if not salida.exists():
        candidatos = list(destino.glob("*.xlsx"))
        if not candidatos:
            raise RuntimeError("LibreOffice no produjo ningun archivo")
        salida = candidatos[0]
    return salida, destino


def columnas_de(hoja_plan):
    """De nombre interno de campo a letra de columna, para los campos que
    declara la descripcion. Las columnas que anade la arquitectura (las de
    clave, que empiezan por "__") no estan en el oraculo."""
    columnas = {}
    for enc in hoja_plan.get("encabezados") or []:
        if enc["campo"].startswith("__"):
            continue
        letra = "".join(c for c in enc["celda"] if c.isalpha())
        columnas[enc["campo"]] = letra
    return columnas


def valores_de_fila(envoltura):
    """Los valores de una fila del oraculo, que viene como [valores, marcas]
    o como el diccionario de valores a secas."""
    return envoltura[0] if isinstance(envoltura, list) else envoltura


def verificar_coleccion(hoja_plan, hoja, filas_esperadas):
    """Compara la hoja de datos de una coleccion con sus filas del oraculo.

    Devuelve (comprobadas, fallos)."""
    columnas = columnas_de(hoja_plan)
    comprobadas = 0
    fallos = []
    for i, envoltura in enumerate(filas_esperadas):
        fila_valores = valores_de_fila(envoltura)
        fila = hoja_plan["fila_primera"] + i
        for campo, letra in columnas.items():
            if campo not in fila_valores:
                continue
            real = normalizar(hoja["{}{}".format(letra, fila)].value)
            esperado_v = normalizar(fila_valores[campo])
            comprobadas += 1
            if not iguales(real, esperado_v):
                fallos.append("{}!{}{}  esperado {!r}  obtenido {!r}".format(
                    hoja.title, letra, fila, esperado_v, real))
    # La columna de estado: lo que dicen las marcas que disparan, en palabras.
    estado = hoja_plan.get("estado")
    if estado:
        for i, envoltura in enumerate(filas_esperadas):
            if not isinstance(envoltura, list) or len(envoltura) < 3:
                continue
            fila = hoja_plan["fila_primera"] + i
            real = normalizar(hoja["{}{}".format(estado, fila)].value)
            esperado_v = normalizar(envoltura[2])
            comprobadas += 1
            if not iguales(real, esperado_v):
                fallos.append("{}!{}{}  estado esperado {!r}  obtenido {!r}".format(
                    hoja.title, estado, fila, esperado_v, real))
    return comprobadas, fallos


def verificar(plano_path, esperado_path, xlsx_path, cruces_path=None):
    """Compara el libro recalculado con los oraculos.

    Devuelve (comprobadas, fallos, por_hoja), donde POR_HOJA es una lista de
    (nombre de hoja, celdas comparadas)."""
    plano = json.loads(Path(plano_path).read_text(encoding="utf-8"))
    esperado = json.loads(Path(esperado_path).read_text(encoding="utf-8"))
    cruces = (json.loads(Path(cruces_path).read_text(encoding="utf-8"))
              if cruces_path else None)

    recalculado, temporal = recalcular_con_libreoffice(xlsx_path)
    try:
        libro = load_workbook(recalculado, data_only=True)
        fallos = []
        por_hoja = []
        colecciones_vistas = set()

        for hoja_plan in plano["hojas"]:
            nombre = hoja_plan["nombre"]
            if nombre not in libro.sheetnames:
                fallos.append("falta la hoja {!r}".format(nombre))
                continue
            hoja = libro[nombre]

            if "coleccion" in hoja_plan:
                coleccion = hoja_plan["coleccion"]
                colecciones_vistas.add(coleccion)
                if coleccion not in esperado:
                    fallos.append("{}: el oraculo no trae la coleccion {!r}".format(
                        nombre, coleccion))
                    continue
                n, f = verificar_coleccion(hoja_plan, hoja, esperado[coleccion])
                if esperado[coleccion] and n == 0:
                    f.append("{}: no se comparo ninguna celda".format(nombre))
            elif "relacion" in hoja_plan:
                n, f = verificar_relacion(hoja_plan, hoja, esperado)
            elif "cruce" in hoja_plan:
                # Las tablas cruzadas se comprueban todas juntas mas abajo.
                continue
            elif "auxiliar" in hoja_plan:
                por_hoja.append((nombre, "auxiliar"))
                continue
            else:
                fallos.append("{}: la hoja no dice que es".format(nombre))
                continue
            por_hoja.append((nombre, n))
            fallos.extend(f)

        for coleccion in esperado:
            if coleccion not in colecciones_vistas:
                fallos.append("la coleccion {!r} no tiene hoja en el libro".format(
                    coleccion))

        if cruces is not None:
            for nombre, n, f in verificar_cruces(plano, cruces, libro, normalizar):
                por_hoja.append((nombre, n))
                fallos.extend(f)
        elif any("cruce" in h for h in plano["hojas"]):
            fallos.append("el libro tiene tablas cruzadas y no se paso el oraculo de cruces")

        fallos.extend("error de formula en " + e for e in errores_del_libro(libro))
        comprobadas = sum(n for _, n in por_hoja if isinstance(n, int))
        return comprobadas, fallos, por_hoja
    finally:
        shutil.rmtree(temporal, ignore_errors=True)


def main(argv):
    if len(argv) not in (4, 5):
        print(__doc__)
        return 2
    comprobadas, fallos, por_hoja = verificar(*argv[1:])
    nombre = Path(argv[3]).name
    detalle = "\n".join("        {:<34} {}".format(h, n) for h, n in por_hoja)
    if fallos:
        print("FALLA  {}: {} celdas comprobadas, {} discrepancias"
              .format(nombre, comprobadas, len(fallos)))
        for f in fallos[:25]:
            print("   " + f)
        print(detalle)
        return 1
    print("  ok   {}: {} celdas evaluadas por LibreOffice coinciden con el "
          "evaluador de referencia".format(nombre, comprobadas))
    print(detalle)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
