"""Comprueba que una fecha tecleada de verdad compara igual que una escrita
como texto.

EL PROBLEMA QUE RESUELVE

El generador escribe las fechas como texto ISO ("2026-10-12"). Pero quien usa
el libro teclea una fecha, y la hoja de calculo la guarda como un numero de
serie. Antes, una comparacion como "entrego despues del 10 de octubre"
comparaba ese numero con el texto "2026-10-10" y daba FALSO, porque para la
hoja de calculo un numero siempre es menor que un texto. La auditoria del
2026-10-17 lo midio en el caso de las notas.

Esta sonda escribe una fecha de verdad en una celda, recalcula con
LibreOffice y comprueba otra celda. La usa ./demostrar-robustez.sh.

Uso:
    python3 sonda_fecha_tecleada.py <libro.xlsx> <hoja> <celda> <AAAA-MM-DD>
                                    <celda-resultado> <valor-esperado>
"""

import datetime
import shutil
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from verificar import recalcular_con_libreoffice


def main(argv):
    if len(argv) != 7:
        print(__doc__)
        return 2
    libro, hoja, celda, fecha, celda_resultado, esperado = argv[1:]
    temporal = Path(tempfile.mkdtemp(prefix="situacion-sonda-"))
    try:
        copia = temporal / Path(libro).name
        wb = load_workbook(libro)
        wb[hoja][celda].value = datetime.date.fromisoformat(fecha)
        wb.save(copia)
        recalculado, otro = recalcular_con_libreoffice(copia)
        try:
            valor = load_workbook(recalculado, data_only=True)[hoja][celda_resultado].value
        finally:
            shutil.rmtree(otro, ignore_errors=True)
    finally:
        shutil.rmtree(temporal, ignore_errors=True)
    nombre = "{}!{} con {} tecleada en {}".format(hoja, celda_resultado, fecha, celda)
    if str(valor) != esperado:
        print("FALLA  {}: esperado {!r}, obtenido {!r}".format(nombre, esperado, valor))
        return 1
    print("  ok   {}: {!r}".format(nombre, valor))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
