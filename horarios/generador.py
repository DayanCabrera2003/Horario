"""Orquestador: carga config → construye el workbook → guarda el .xlsx."""
import datetime
from pathlib import Path

from openpyxl import Workbook

from comun.portada import (
    NOMBRE_HOJA as HOJA_PORTADA, SE_ESCRIBE, SE_CALCULA, SON_DATOS,
    construir_portada,
)
from horarios.config import cargar_facultad, cargar_horarios
from horarios.hoja_datos import construir_hoja_datos
from horarios.hoja_grupo import construir_hoja_grupo
from horarios.hoja_ocupacion import NOMBRE_HOJA as HOJA_AULAS, construir_hoja_ocupacion
from horarios.hoja_datos import NOMBRE_HOJA as HOJA_AUXILIAR
from horarios.hoja_aulas import (
    NOMBRE_HOJA as HOJA_LISTA_AULAS, construir_hoja_aulas)
from horarios.hoja_asignaturas import (
    NOMBRE_HOJA as HOJA_ASIGNATURAS, construir_hoja_asignaturas)
from horarios.hoja_listado_grupos import (
    NOMBRE_HOJA as HOJA_GRUPOS, construir_hoja_grupos)
from horarios.hoja_estructura import (
    NOMBRE_HOJA as HOJA_ESTRUCTURA, construir_hoja_estructura)
from horarios.hoja_profesores import (
    NOMBRE_HOJA as HOJA_PROFESORES, construir_hoja_profesores)
from horarios.hoja_docencia import (
    NOMBRE_HOJA as HOJA_DOCENCIA, construir_hoja_docencia)


def _indice(hojas_grupo, con_profesores: bool) -> tuple:
    """Que hay en cada hoja y si se escribe a mano, para el indice de la portada.

    `hojas_grupo` son pares (nombre_de_hoja, grupo) tomados de las hojas ya
    creadas, no una copia de los ids. `Datos` no aparece: esta oculta y solo es
    fontaneria de formulas.
    """
    # Las dos hojas de profesorado solo estan si el YAML las declara, asi que
    # el indice tiene que preguntarlo en vez de darlas por hechas.
    profesorado = (
        (HOJA_PROFESORES, "Los profesores de la facultad", SON_DATOS),
        (HOJA_DOCENCIA, "Quién imparte cada asignatura en cada grupo", SON_DATOS),
    ) if con_profesores else ()
    return (
        (HOJA_LISTA_AULAS, "Las aulas de la facultad", SON_DATOS),
        (HOJA_ASIGNATURAS, "Qué se imparte en cada año y con qué frecuencia", SON_DATOS),
        (HOJA_GRUPOS, "Los grupos y de dónde sale el id de cada uno", SON_DATOS),
        (HOJA_ESTRUCTURA, "Días, turnos y tamaño del libro", SON_DATOS),
        *profesorado,
        (HOJA_AULAS, "Qué aula está ocupada en cada día y turno", SE_CALCULA),
        *((nombre, f"Carrera {grupo.carrera} · año {grupo.anio}", SE_ESCRIBE)
          for nombre, grupo in hojas_grupo),
    )


def _subtitulo(facultad) -> str:
    """Resumen de tamano del libro: el dominio de horarios no tiene nombre de
    facultad ni curso que poner aqui, asi que se dice cuanto abarca."""
    grupos, aulas = len(facultad.grupos), len(facultad.aulas)
    return (f"{grupos} grupo{'s' if grupos != 1 else ''} · "
            f"{aulas} aula{'s' if aulas != 1 else ''}")


def generar(
    config_path: Path,
    horarios_path: Path | None,
    salida: Path,
    generado: datetime.datetime | None = None,
) -> Path:
    """Carga la configuración, construye el workbook (Portada + Datos + grupos +
    Aulas) y guarda en salida.

    `generado` es la fecha que muestra la portada; se puede fijar desde los tests
    para que la salida no dependa del reloj.
    """
    generado = generado or datetime.datetime.now()
    facultad = cargar_facultad(config_path)
    horarios = cargar_horarios(horarios_path, facultad)

    wb = Workbook()
    # Elimina la hoja por defecto; se crearán las nuestras.
    wb.remove(wb.active)

    # Hoja auxiliar primero (define AulasValidas y firmas que otras hojas usan).
    firmas = construir_hoja_datos(wb, facultad)

    # Los datos del problema, delante de todo lo que se deriva de ellos.
    construir_hoja_aulas(wb, facultad)
    construir_hoja_asignaturas(wb, facultad)
    construir_hoja_grupos(wb, facultad)
    construir_hoja_estructura(wb, facultad)
    # Solo si el YAML declara profesores: las dos secciones son opcionales.
    construir_hoja_profesores(wb, facultad)
    construir_hoja_docencia(wb, facultad)

    # La ocupacion, antes de las hojas de grupo: es la vista de conjunto.
    construir_hoja_ocupacion(wb, facultad, firmas)

    # Una hoja por grupo.
    hojas_grupo = []
    for grupo in facultad.grupos:
        ws = wb.create_sheet(grupo.id)
        construir_hoja_grupo(ws, grupo, facultad, horario=horarios.get(grupo.id))
        hojas_grupo.append((ws.title, grupo))

    # La hoja auxiliar, al final del todo: es fontanería y va oculta.
    wb.move_sheet(HOJA_AUXILIAR, offset=len(wb.sheetnames) - 1)

    # La portada va la última porque se inserta en el índice 0: así ya sabe qué
    # hojas existen y queda delante de todas.
    construir_portada(
        wb,
        titulo="Horario de la facultad",
        subtitulo=_subtitulo(facultad),
        origen=str(config_path),
        hojas=_indice(hojas_grupo, bool(facultad.profesores)),
        generado=generado,
    )
    # El libro abre por la portada: si abriera por otra hoja nadie la leería.
    wb.active = wb[HOJA_PORTADA]

    wb.save(salida)
    return salida
