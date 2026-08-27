"""Proteccion de hojas: bloquear lo que se calcula solo, dejar editable lo que
se llena a mano.

En OOXML toda celda nace con `locked=True`, pero ese atributo no hace nada
mientras la hoja no este protegida. Al proteger, se cierra todo de golpe; por
eso el desbloqueo de lo editable tiene que ser explicito.

Sin contrasena a proposito: el objetivo es evitar que se borre una formula sin
querer, no proteger secretos. Quien necesite tocar una celda calculada quita la
proteccion desde el menu de Excel o de Calc.
"""
from openpyxl.styles import Protection

_DESBLOQUEADA = Protection(locked=False)


def celda_unica(ref: str) -> str:
    """Convierte una referencia de celda suelta en un rango de una celda.

    `ws["F4"]` devuelve una Cell y no una tupla de filas, asi que desbloquear
    una celda concreta exige pasarla como "F4:F4". Este helper evita que cada
    consumidor repita esa duplicacion a mano.
    """
    return f"{ref}:{ref}"


def desbloquear(ws, rangos) -> None:
    """Marca como editables las celdas de `rangos`.

    Cada entrada puede ser un rango simple ("B2:C3") o varios separados por
    espacios ("B4:F4 B6:F6"), que es el formato que devuelven las funciones de
    `layout` para el formato condicional. `ws[...]` no acepta el segundo, asi
    que se parte aqui.

    Una celda suelta hay que pasarla con el rango duplicado ("F4:F4"): "F4"
    sin dos puntos devuelve una `Cell` en vez de una tupla de filas, y el
    bucle revienta con `TypeError: 'Cell' object is not iterable`.

    Se aplica celda a celda porque en openpyxl la proteccion es un atributo de
    estilo por celda, no de rango. Asignarla no toca fuente, relleno ni borde:
    cada atributo de estilo es independiente.
    """
    for entrada in rangos:
        for rango in str(entrada).split():
            for fila in ws[rango]:
                for celda in fila:
                    celda.protection = _DESBLOQUEADA


def proteger_hoja(ws, editables=(), permitir_orden: bool = False,
                  permitir_filtro: bool = False) -> None:
    """Protege `ws` dejando editables los rangos de `editables`.

    `permitir_orden` habilita ordenar y `permitir_filtro` habilita autofiltrar.
    Al proteger, OOXML prohibe las dos por defecto, y en las hojas de consulta
    eso estorba mas de lo que ayuda. Cada atributo de `SheetProtection`
    significa "esta operacion queda prohibida", de ahi que habilitarlas sea
    ponerlas en False.

    Son dos permisos separados a proposito: filtrar solo esconde filas, pero
    ordenar las mueve de sitio. En una hoja cuyas filas tienen una gemela por
    posicion en otra hoja, ordenar descuadraria la gemela sin avisar, asi que
    ahi se concede el filtro y se niega el orden.
    """
    desbloquear(ws, editables)
    ws.protection.sheet = True
    if permitir_orden:
        ws.protection.sort = False
    if permitir_filtro:
        ws.protection.autoFilter = False
