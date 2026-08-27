"""Constructor generico de hoja de listado: una tabla simple con el formato de
la casa.

Las hojas de listado son las que sacan a la luz los datos del problema (quienes
son los profesores, que aulas hay, que asignaturas existen). Todas tienen la
misma forma -encabezados arriba y una fila por dato-, asi que se construyen
desde aqui en vez de repetir el mismo codigo en cada generador.

No decide que datos van: recibe los encabezados y las filas ya resueltos, igual
que `leyenda.py` recibe los colores ya elegidos.
"""
from openpyxl.utils import get_column_letter

from comun import formato, proteccion, vista
from comun import estilos_base as estilos

# Los encabezados van en la primera fila, sin fila de titulo: asi el congelado
# es "A2" y la tabla empieza donde la vista espera. Las hojas de reporte si
# llevan titulo, pero un listado se lee como una tabla, no como un informe.
FILA_ENCABEZADO = 1
FILA_PRIMER_DATO = FILA_ENCABEZADO + 1

# Mismo alto que las tablas de los reportes: parte del padding aproximado para
# Calc, junto con la sangria y el centrado vertical.
ALTO_FILA = 22


def construir_hoja_listado(wb, nombre: str, encabezados, filas,
                           color_encabezado: str):
    """Crea la hoja `nombre` con `encabezados` y una fila por elemento de
    `filas`, y la devuelve.

    `filas` es un iterable de secuencias con tantos valores como encabezados.
    Puede venir vacio: la hoja sale igual, solo con los encabezados.
    """
    ws = wb.create_sheet(nombre)
    filas = list(filas)

    for i, texto in enumerate(encabezados):
        ws.cell(row=FILA_ENCABEZADO, column=i + 1, value=texto)
    for f, valores in enumerate(filas):
        for c, valor in enumerate(valores):
            ws.cell(row=FILA_PRIMER_DATO + f, column=c + 1, value=valor)

    _aplicar_presentacion(ws, len(encabezados), len(filas), color_encabezado)

    # Encabezados a la vista al bajar por una lista larga.
    ws.freeze_panes = f"A{FILA_PRIMER_DATO}"
    # Un listado es de consulta: nada editable, pero se deja ordenar y filtrar
    # para buscar dentro de el. En la fase 3a estas hojas pasan a ser la fuente
    # de los desplegables y se desbloquean.
    proteccion.proteger_hoja(ws, permitir_orden=True, permitir_filtro=True)
    return ws


def _aplicar_presentacion(ws, n_cols: int, n_filas: int,
                          color_encabezado: str) -> None:
    columnas = [get_column_letter(i + 1) for i in range(n_cols)]
    ultima_fila = FILA_ENCABEZADO + n_filas
    rango = f"A{FILA_ENCABEZADO}:{columnas[-1]}{ultima_fila}"

    formato.aplicar_estilo_encabezado(
        ws, [f"{col}{FILA_ENCABEZADO}" for col in columnas],
        estilos.fuente_encabezado(), estilos.fill(color_encabezado))
    formato.aplicar_borde_tabla(ws, rango, interno=estilos.lado_fino(),
                                externo=estilos.lado_medio())
    formato.aplicar_alineacion(ws, rango, estilos.alineacion_padding())
    formato.aplicar_alto_filas(ws, FILA_ENCABEZADO, ultima_fila, ALTO_FILA)
    formato.autoajustar_columnas(ws)
    # Hoja de tabla: los bordes ya delimitan; la cuadricula de fondo compite.
    vista.ocultar_cuadricula(ws)
