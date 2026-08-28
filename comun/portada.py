"""Hoja de portada: la primera de cada libro.

Dice de donde salio el libro, cuando, que hay en cada hoja y que hacer para
cambiar los datos. Existe porque el tutor tuvo que preguntar donde se entraban
los datos de la docencia: el libro no lo decia por ninguna parte.

La fecha de generacion se inyecta en vez de leerse del reloj para que la salida
sea reproducible y los tests puedan fijarla.
"""
import datetime

from openpyxl.utils import quote_sheetname
from openpyxl.worksheet.hyperlink import Hyperlink

from comun import formato, proteccion, vista
from comun import estilos_base as estilos

NOMBRE_HOJA = "Portada"

_FILA_TITULO = 2
_FILA_ORIGEN = 4
_FILA_INSTRUCCIONES = 7
_FILA_INDICE = 11

# Etiquetas del indice. Son tres y no dos porque una hoja de listado no es
# ninguna de las dos cosas de antes: no se calcula sola, pero tampoco es donde se
# planifica. Es donde estan los datos del problema, que es exactamente lo que el
# tutor no encontraba.
SE_ESCRIBE = "se escribe"
SE_CALCULA = "se calcula"
SON_DATOS = "datos del problema"

_INSTRUCCIONES = (
    "Las casillas de fondo blanco se escriben a mano. Las demás se calculan "
    "solas y están bloqueadas: si necesitas tocarlas, quita la protección de "
    "la hoja desde el menú de Excel o de Calc."
)

# Aviso de las hojas de datos. Va en la portada y no solo en la guia porque el
# efecto de dejar un hueco es silencioso: la lista se corta ahi y lo que queda
# debajo desaparece de los desplegables sin ningun mensaje. Comprobado en Calc.
_AVISO_LISTAS = (
    "En las hojas de datos puedes añadir filas al final de la lista, en las "
    "líneas libres que ya vienen preparadas. No dejes filas en blanco en medio: "
    "la lista se corta ahí y lo que quede debajo desaparece de los desplegables."
)


def _enlazar_a_hoja(celda, nombre: str) -> None:
    """Convierte `celda` en un salto a la celda A1 de la hoja `nombre`.

    Se construye un Hyperlink con `location` en vez de asignar la cadena
    "#'Hoja'!A1" a `celda.hyperlink`: esa forma abreviada la guarda openpyxl
    como relacion externa (TargetMode="External"), que no es lo que es un salto
    dentro del mismo libro. `location` produce el enlace interno del formato.
    """
    celda.hyperlink = Hyperlink(ref=celda.coordinate,
                                location=f"{quote_sheetname(nombre)}!A1")


def _encajar_parrafos(ws, filas) -> None:
    """Deja los parrafos largos dentro de la pagina.

    Son frases de 200 caracteres en una celda suelta: sin combinar ni ajustar el
    texto se salen de la hoja, y al imprimir (o al exportar a PDF) lo que no cabe
    se pierde y la portada se parte en dos paginas. Se combinan las tres primeras
    columnas y se activa el ajuste de linea, con el alto necesario para las
    lineas que salgan.
    """
    ancho = sum(ws.column_dimensions[c].width or 10 for c in ("A", "B", "C"))
    for fila in filas:
        celda = ws[f"A{fila}"]
        if not celda.value:
            continue
        ws.merge_cells(f"A{fila}:C{fila}")
        celda.alignment = estilos.alineacion_ajuste()
        lineas = max(1, -(-len(str(celda.value)) // max(20, int(ancho))))
        ws.row_dimensions[fila].height = 15 * lineas


def construir_portada(wb, titulo: str, subtitulo: str, origen: str,
                      hojas, generado: datetime.datetime,
                      instrucciones: str = _INSTRUCCIONES) -> None:
    """Crea la hoja Portada como primera del libro.

    `hojas` es un iterable de (nombre, descripcion, etiqueta) en el orden en que
    se quieren listar; `etiqueta` es una de SE_ESCRIBE, SE_CALCULA o SON_DATOS.
    `generado` se pasa siempre de forma explicita.
    """
    ws = wb.create_sheet(NOMBRE_HOJA, index=0)

    ws[f"A{_FILA_TITULO}"] = titulo
    ws[f"A{_FILA_TITULO + 1}"] = subtitulo

    ws[f"A{_FILA_ORIGEN}"] = "Generado desde"
    ws[f"B{_FILA_ORIGEN}"] = origen
    ws[f"A{_FILA_ORIGEN + 1}"] = "Fecha"
    ws[f"B{_FILA_ORIGEN + 1}"] = generado.strftime("%d/%m/%Y %H:%M")

    ws[f"A{_FILA_INSTRUCCIONES}"] = "Donde se cambian los datos"
    ws[f"A{_FILA_INSTRUCCIONES + 1}"] = instrucciones
    ws[f"A{_FILA_INSTRUCCIONES + 2}"] = _AVISO_LISTAS

    ws[f"A{_FILA_INDICE}"] = "Las hojas de este libro"
    for i, (nombre, descripcion, etiqueta) in enumerate(hojas):
        fila = _FILA_INDICE + 1 + i
        celda = ws[f"A{fila}"]
        celda.value = nombre
        _enlazar_a_hoja(celda, nombre)
        ws[f"B{fila}"] = descripcion
        ws[f"C{fila}"] = etiqueta

    formato.autoajustar_columnas(ws, extra=4)
    # La columna A la fija el texto mas largo, y los dos parrafos de arriba son
    # de 200 caracteres: sin esto se lleva el ancho maximo y empuja el indice
    # fuera de la pagina. Se dimensiona con lo que de verdad vive en ella.
    formato.fijar_ancho_por_textos(
        ws, "A", [nombre for nombre, _, _ in hojas] + ["Generado desde",
                                                      "Las hojas de este libro"],
        extra=4)
    _encajar_parrafos(ws, (_FILA_INSTRUCCIONES + 1, _FILA_INSTRUCCIONES + 2))
    vista.ocultar_cuadricula(ws)
    # La portada no tiene nada que editar.
    proteccion.proteger_hoja(ws)
