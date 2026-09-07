"""Hoja `Carga por profesor`: reporte de carga por profesor, calculado con formulas.

Un bloque por hueco del claustro -los profesores declarados y su reserva-, con
sus datos (id, nombre, grado, tope efectivo), el detalle de lo que imparte y
una fila TOTAL. El id del bloque tampoco se escribe al generar: se lee de la
hoja `Profesores`, de modo que un profesor anadido alli tiene su bloque y
cambiar un id no deja el bloque buscando a quien ya no existe.

El detalle no se conoce al generar (depende de lo que se elija en Asignacion),
asi que se reservan
`filas_por_profesor` lineas rellenadas con BUSCARV sobre la tabla auxiliar
CargaPorProfesor de la hoja Datos (claves '<id>#<n>'). Son formulas planas: se
comportan igual en Excel y en Calc. La ultima linea reservada avisa '(+N más)'
si el profesor imparte mas cosas de las que caben en el bloque.
"""
from openpyxl.utils import quote_sheetname

from comun import formato, leyenda
from comun import proteccion
from comun import vista
from comun.hoja_listado import FILA_PRIMER_DATO
from departamento import estilos
from departamento import layout as L
from departamento.hoja_asignacion import NOMBRE_HOJA as HOJA_ASIGNACION
from departamento.hoja_profesores import (
    NOMBRE_HOJA as HOJA_CLAUSTRO, COL_ID as COL_CLAUSTRO_ID)
from departamento.modelo import Departamento

# El nombre dice lo que la hoja es: un reporte de carga, no el listado del
# claustro (ese es la hoja "Profesores", que si son los datos del problema).
NOMBRE_HOJA = "Carga por profesor"

# Columnas del detalle: BUSCARV sobre CargaPorProfesor (clave en la col. 1).
# A=asignatura (col 2), B=tipo (3), C=grupo (4), D=horas (5).
_COLS_DETALLE = (("A", 2), ("B", 3), ("C", 4), ("D", 5))


def construir_hoja_carga(wb, depto: Departamento) -> None:
    ws = wb.create_sheet(NOMBRE_HOJA)
    ws[f"A{L.FILA_TITULO}"] = f"{NOMBRE_HOJA} — {depto.nombre} — {depto.semestre}"
    ws[f"A{L.FILA_TITULO}"].font = estilos.fuente_encabezado()

    fpp = depto.filas_por_profesor
    # Un bloque por hueco del claustro, no por profesor declarado: la hoja
    # `Profesores` tiene filas libres y quien se anada alli necesita el suyo.
    huecos = depto.capacidad_profesores()
    for i in range(huecos):
        _construir_bloque(ws, depto, i)

    ultimo_total = L.prof_fila_total(huecos - 1, fpp)
    # Padding aproximado para Calc: filas mas altas en toda la zona de bloques.
    formato.aplicar_alto_filas(ws, L.PROF_FILA_PRIMER_BLOQUE, ultimo_total,
                               estilos.ALTO_FILA)
    _escribir_leyenda(ws, ultimo_total + 2)
    ws.freeze_panes = f"A{L.PROF_FILA_PRIMER_BLOQUE}"
    formato.autoajustar_columnas(ws, extra=4)
    # La columna A del detalle muestra nombres de asignatura por formula
    # (autoajustar la ignora): el ancho se fija con los nombres posibles.
    formato.fijar_ancho_por_textos(
        ws, "A", [a.nombre for a in depto.asignaturas] + ["Asignatura"], extra=4)

    # Hoja de solo lectura: el tope se edita en el claustro (hoja `Profesores`)
    # y aqui llega por BUSCARV. Era la arruga declarada de la fase 1.
    proteccion.proteger_hoja(ws)
    # Pestana de navegacion: salvo los topes, todo sale de Asignacion por formula.
    vista.colorear_pestana(ws, estilos.COLOR_PESTANA_CALCULO)
    # Hoja de reporte: las tablas ya van bordeadas y la cuadricula de fondo
    # compite con esos bordes.
    vista.ocultar_cuadricula(ws)


def _rango_asignacion(col: str, n_filas: int) -> str:
    hoja = quote_sheetname(HOJA_ASIGNACION)
    return (f"{hoja}!${col}${L.FILA_PRIMERA_CARGA}"
            f":${col}${L.fila_carga(n_filas - 1)}")


def _construir_bloque(ws, depto: Departamento, idx: int) -> None:
    fpp = depto.filas_por_profesor
    # La reserva de Asignacion entra en el rango: si no, las filas de carga
    # creadas a mano no sumarian para nadie. Es un fallo silencioso -no da
    # error, solo horas de menos- y por eso va explicito.
    n_filas = depto.capacidad_filas()
    rango_prof = _rango_asignacion(L.COL_PROFESOR, n_filas)

    # Cabecera y valores del profesor.
    fila_cab = L.prof_fila_cabecera(idx, fpp)
    _fila_encabezado(ws, fila_cab, ("Id", "Nombre", "Grado", "Tope horas"))
    fila_val = L.prof_fila_valores(idx, fpp)
    # Ni el id se escribe al generar: sale del hueco gemelo del claustro, que es
    # donde se anaden y se corrigen los profesores. Todo lo demas del bloque se
    # cuelga de esta celda, asi que un bloque libre queda entero en blanco.
    id_bloque = f"$A${fila_val}"
    fila_claustro = FILA_PRIMER_DATO + idx
    celda_claustro = f"{HOJA_CLAUSTRO}!{COL_CLAUSTRO_ID}{fila_claustro}"
    ws[f"A{fila_val}"] = (
        f'=IF({celda_claustro}="","",{celda_claustro})')
    conteo = f"COUNTIF({rango_prof},{id_bloque})"
    for col, columna_tabla in (("B", 2), ("C", 3), ("D", 4)):
        ws[f"{col}{fila_val}"] = (
            f'=IF({id_bloque}="","",'
            f'IFERROR(VLOOKUP({id_bloque},ProfesoresTabla,{columna_tabla},0),""))')

    # Subcabecera y detalle reservado.
    _fila_encabezado(ws, L.prof_fila_subcabecera(idx, fpp),
                     ("Asignatura", "Tipo", "Grupo", "Horas"))
    for k in range(fpp):
        fila = L.prof_fila_detalle(idx, k, fpp)
        for col, col_tabla in _COLS_DETALLE:
            clave = f'{id_bloque}&"#{k + 1}"'
            buscar = (f"IFERROR(VLOOKUP({clave},"
                      f'CargaPorProfesor,{col_tabla},0),"")')
            if col == "A" and k == fpp - 1:
                # Ultima linea reservada: si hay desborde, avisa cuantas faltan
                # (la propia linea deja de mostrarse, por eso se suma 1).
                buscar = (f'IF({conteo}>{fpp},'
                          f'"(+"&({conteo}-{fpp}+1)&" más)",{buscar})')
            ws[f"{col}{fila}"] = f'=IF({id_bloque}="","",{buscar})'

    # TOTAL de horas del profesor y alerta de sobrecarga.
    fila_total = L.prof_fila_total(idx, fpp)
    ws[f"A{fila_total}"] = "TOTAL"
    ws[f"A{fila_total}"].font = estilos.fuente_encabezado()
    ws[f"D{fila_total}"] = (
        f'=IF({id_bloque}="","",SUMIF({rango_prof},{id_bloque},'
        f"{_rango_asignacion(L.COL_HORAS, n_filas)}))")
    # Rojo en la fila TOTAL cuando el acumulado supera el tope del bloque. La
    # regla se crea siempre, tenga tope o no al generar: el tope se edita en el
    # claustro, asi que ponerla solo a quien ya lo traia dejaria sin alerta
    # justo a quien se lo pongan despues. La guarda del <> "" evita que dispare
    # con la casilla vacia, donde no hay nada con que comparar.
    ws.conditional_formatting.add(
        f"A{fila_total}:D{fila_total}",
        estilos.regla_formula(
            f'AND($D${fila_val}<>"",$D${fila_total}>$D${fila_val})',
            estilos.COLOR_SOBRECARGA))

    # Todas las horas del bloque son enteras: el tope, las lineas de detalle y
    # el TOTAL. Sin formato explicito, Calc las mostraria segun la configuracion
    # regional de quien abra el libro. El tope va suelto porque entre el y el
    # detalle quedan las dos filas de encabezado, que son texto.
    formato.aplicar_formato_numero(ws, f"D{fila_val}")
    formato.aplicar_formato_numero(
        ws, f"D{L.prof_fila_detalle(idx, 0, fpp)}:D{fila_total}")

    rango = f"A{fila_cab}:D{fila_total}"
    formato.aplicar_borde_tabla(ws, rango, interno=estilos.lado_fino(),
                                externo=estilos.lado_medio())
    formato.aplicar_alineacion(ws, rango, estilos.alineacion_padding())


def _fila_encabezado(ws, fila: int, textos) -> None:
    for i, texto in enumerate(textos):
        celda = ws.cell(row=fila, column=i + 1, value=texto)
        celda.font = estilos.fuente_encabezado()
        celda.fill = estilos.fill(estilos.COLOR_ENCABEZADO)


def _escribir_leyenda(ws, fila: int) -> None:
    leyenda.escribir_leyenda(ws, f"A{fila}", (
        (estilos.COLOR_SOBRECARGA, "Profesor por encima de su tope de horas"),
    ))
