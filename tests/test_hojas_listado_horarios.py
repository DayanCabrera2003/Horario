"""Tests de las hojas de listado del generador de horarios.

Aqui la estructura del problema estaba casi toda implicita: las aulas vivian
escondidas en la hoja auxiliar, las asignaturas solo aparecian repetidas en la
tabla de cada grupo, y los ids de grupo no estaban escritos en ninguna parte
(salian de combinar carreras x anos x sesiones x grupos).
"""
from openpyxl import Workbook

from horarios.hoja_aulas import construir_hoja_aulas
from horarios.hoja_asignaturas import construir_hoja_asignaturas
from horarios.hoja_listado_grupos import construir_hoja_grupos
from horarios.hoja_estructura import construir_hoja_estructura
from horarios.modelo import Asignatura, Anio, Grupo, Facultad

FACULTAD = Facultad(
    aulas=("Aula 1", "Aula 2", "Lab"),
    dias=("Lunes", "Martes", "Miércoles"),
    turnos=6,
    grupos=(Grupo(carrera="C", anio=1, sesion=1, numero=1),
            Grupo(carrera="C", anio=1, sesion=1, numero=2),
            Grupo(carrera="M", anio=2, sesion=1, numero=1)),
    anios={
        "C1": Anio(carrera="C", numero=1, asignaturas=(
            Asignatura(id="AMI-C", nombre="Análisis Matemático I (Conf)", frecuencia=1),
            Asignatura(id="AMI-CP", nombre="Análisis Matemático I (C.P.)", frecuencia=2))),
        "M2": Anio(carrera="M", numero=2, asignaturas=(
            Asignatura(id="GA", nombre="Geometría Analítica", frecuencia=3),)),
    },
)


def _hoja(constructor):
    wb = Workbook()
    wb.remove(wb.active)
    constructor(wb, FACULTAD)
    return wb[wb.sheetnames[0]]


def _columna(ws, letra):
    return [ws[f"{letra}{f}"].value for f in range(2, ws.max_row + 1)
            if ws[f"{letra}{f}"].value is not None]


def test_aulas_saca_a_la_luz_la_lista_maestra():
    ws = _hoja(construir_hoja_aulas)
    assert ws.title == "Aulas"
    assert ws["A1"].value == "Aula"
    assert _columna(ws, "A") == ["Aula 1", "Aula 2", "Lab"]


def test_asignaturas_aplana_los_anios():
    ws = _hoja(construir_hoja_asignaturas)
    assert ws.title == "Asignaturas"
    assert [c.value for c in ws[1][:5]] == ["Carrera", "Año", "Id", "Nombre",
                                            "Frecuencia"]
    # Una fila por (ano, asignatura): tres en total.
    assert _columna(ws, "C") == ["AMI-C", "AMI-CP", "GA"]
    assert ws["A2"].value == "C"
    assert ws["B2"].value == 1
    assert ws["E2"].value == 1


def test_grupos_escribe_los_ids_que_hoy_hay_que_deducir():
    ws = _hoja(construir_hoja_grupos)
    assert ws.title == "Grupos"
    assert [c.value for c in ws[1][:5]] == ["Grupo", "Carrera", "Año",
                                            "Sesión", "Número"]
    assert _columna(ws, "A") == ["C111", "C112", "M211"]


def test_estructura_dice_los_dias_y_los_turnos():
    ws = _hoja(construir_hoja_estructura)
    assert ws.title == "Estructura"
    assert [c.value for c in ws[1][:2]] == ["Concepto", "Valor"]
    conceptos = dict(zip(_columna(ws, "A"), _columna(ws, "B")))
    assert conceptos["Días"] == "Lunes, Martes, Miércoles"
    assert conceptos["Turnos por día"] == 6
    assert conceptos["Carreras"] == "C, M"
    assert conceptos["Años"] == "C1, M2"
    assert conceptos["Grupos"] == 3


def test_los_listados_quedan_protegidos_y_ordenables():
    for constructor in (construir_hoja_aulas, construir_hoja_asignaturas,
                        construir_hoja_grupos, construir_hoja_estructura):
        ws = _hoja(constructor)
        assert ws.protection.sheet is True, ws.title
        assert ws.protection.sort is False, ws.title
