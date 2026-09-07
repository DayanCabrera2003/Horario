"""Tests de la hoja `Cobertura por asignatura`.

Es el segundo pedido del tutor de 2026-09-07: que anadir una asignatura en la
hoja Asignaturas se refleje aqui. Por eso la hoja dejo de ser un bloque por
asignatura -geometria decidida al generar- y paso a ser una tabla con una fila
por hueco de la hoja Asignaturas, reserva incluida.
"""
from openpyxl import Workbook

from departamento.modelo import Profesor, Asignatura, Departamento
from departamento.hoja_datos import construir_hoja_datos
from departamento.hoja_asignaturas import construir_hoja_asignaturas
from departamento.hoja_asignacion import construir_hoja_asignacion
from departamento.hoja_cobertura import construir_hoja_cobertura
from departamento import estilos


def _departamento():
    return Departamento(
        nombre="Matemática Aplicada", semestre="2026-2027 / 1",
        tope_horas=160, filas_por_profesor=10,
        filas_carga_reserva=2, asignaturas_reserva=1,
        profesores=(Profesor(id="PIAD", nombre="Pedro I. Alonso", grado="Dr."),),
        asignaturas=(
            Asignatura(id="EST-CC", nombre="Estadística (CC)",
                       carrera="Ciencia de la Computación",
                       horas_conf=32, horas_cp=32, grupos_cp=2),
            Asignatura(id="EST-MAT", nombre="Estadística (Mat)",
                       carrera="Matemática", horas_conf=32, horas_cp=32,
                       grupos_cp=1),
        ),
    )


def _hoja():
    # 2 asignaturas + 1 de reserva -> filas 3, 4 y 5.
    # 5 filas de carga + 2 de reserva -> A4:A10 en la hoja Asignacion.
    wb = Workbook()
    wb.remove(wb.active)
    depto = _departamento()
    construir_hoja_datos(wb, depto)
    construir_hoja_asignaturas(wb, depto)
    construir_hoja_asignacion(wb, depto)
    construir_hoja_cobertura(wb, depto)
    return wb["Cobertura por asignatura"]


def _reglas(ws):
    return [(str(rango.sqref), regla.formula[0],
             regla.dxf.fill.start_color.rgb[-6:])
            for rango, lista in ws.conditional_formatting._cf_rules.items()
            for regla in lista if regla.formula]


def test_encabezados_de_la_tabla():
    ws = _hoja()
    assert [c.value for c in ws["A2":"I2"][0]] == [
        "Id", "Asignatura", "Carrera", "Horas planificadas", "Filas de carga",
        "Horas en filas", "Horas asignadas", "Filas sin profesor", "Estado"]


def test_una_fila_por_hueco_de_la_hoja_asignaturas():
    ws = _hoja()
    assert ws["A3"].value == '=IF(Asignaturas!A2="","",Asignaturas!A2)'
    assert ws["A4"].value == '=IF(Asignaturas!A3="","",Asignaturas!A3)'
    assert ws["A5"].value is not None      # la fila de reserva existe
    assert ws["A6"].value is None          # y no hay ninguna de mas


def test_horas_planificadas_vienen_del_plan():
    # El total que declara la asignatura: Conf + CP por grupo. Si se corrigen
    # sus grupos en la hoja Asignaturas, esta cifra se mueve con ellos.
    assert "Asignaturas!G2" in _hoja()["D3"].value


def test_filas_de_carga_y_horas_creadas():
    ws = _hoja()
    assert ws["E3"].value == (
        '=IF($A3="","",COUNTIF(\'Asignación\'!$A$4:$A$10,$A3))')
    assert ws["F3"].value == (
        '=IF($A3="","",SUMIF(\'Asignación\'!$A$4:$A$10,$A3,'
        "'Asignación'!$F$4:$F$10))")


def test_horas_asignadas_excluyen_las_filas_sin_profesor():
    formula = _hoja()["G3"].value
    assert "SUMIFS" in formula
    assert '"<>"' in formula        # profesor distinto de vacio


def test_filas_sin_profesor():
    formula = _hoja()["H3"].value
    assert "COUNTIFS" in formula
    assert "'Asignación'!$G$4:$G$10" in formula


def test_el_estado_distingue_las_tres_formas_de_faltar():
    formula = _hoja()["I3"].value
    for texto in ("Sin filas de carga", "horas de carga", "profesores",
                  "Completa"):
        assert texto in formula, texto


def test_colores_por_estado():
    reglas = _hoja()
    verdes = [(sq, f) for sq, f, color in _reglas(reglas)
              if color == estilos.COLOR_COMPLETA]
    naranjas = [(sq, f) for sq, f, color in _reglas(reglas)
                if color == estilos.COLOR_INCOMPLETA]
    assert verdes and naranjas
    assert verdes[0][0] == "A3:I5"
    assert '$I3="Completa"' in verdes[0][1]
    # La guarda del id: una fila de reserva no esta incompleta, esta vacia.
    assert '$A3<>""' in naranjas[0][1]


def test_una_asignatura_nueva_aparece_sola():
    # El punto del pedido: la fila de reserva ya trae todas sus formulas, asi
    # que escribir una asignatura en la hoja Asignaturas la hace aparecer aqui
    # sin regenerar nada.
    ws = _hoja()
    for col in "ABCDEFGHI":
        assert str(ws[f"{col}5"].value).startswith("="), col


def test_la_hoja_es_de_solo_lectura():
    # Todo son formulas: las asignaturas se anaden en la hoja Asignaturas y los
    # profesores se eligen en Asignacion.
    ws = _hoja()
    assert ws.protection.sheet is True
    assert ws["A3"].protection.locked is True


def test_la_hoja_lleva_autofiltro_y_congelado():
    ws = _hoja()
    assert ws.auto_filter.ref == "A2:I5"
    assert ws.freeze_panes == "A3"
    # Se puede ordenar: aqui ninguna fila tiene gemela por posicion, cada una
    # se resuelve por su id.
    assert ws.protection.sort is False


def test_las_cifras_son_enteras():
    ws = _hoja()
    for col in ("D", "E", "F", "G", "H"):
        assert ws[f"{col}3"].number_format == "0", col


def test_leyenda_de_colores():
    ws = _hoja()
    textos = [ws[f"B{r}"].value or "" for r in range(6, 11)]
    assert any("ompleta" in t for t in textos), textos


def test_la_hoja_no_muestra_cuadricula():
    assert _hoja().sheet_view.showGridLines is False
