"""Tests del panel de profesores de la hoja Asignacion.

Es el primero de los dos pedidos del tutor de 2026-09-07: al lado de la tabla,
un listado de profesores con el total de horas de sus asignaciones y la
cantidad de asignaturas, calculado solo.
"""
from openpyxl import Workbook

from departamento.modelo import Profesor, Asignatura, Departamento
from departamento.hoja_datos import construir_hoja_datos
from departamento.hoja_asignacion import construir_hoja_asignacion
from departamento import estilos


def _departamento():
    return Departamento(
        nombre="Matemática Aplicada", semestre="2026-2027 / 1",
        tope_horas=160, filas_por_profesor=10,
        filas_carga_reserva=2, profesores_reserva=1,
        profesores=(
            Profesor(id="PIAD", nombre="Pedro I. Alonso", grado="Dr."),
            Profesor(id="MARA", nombre="Maria Ramirez", grado="MSc."),
        ),
        asignaturas=(
            Asignatura(id="EST-CC", nombre="Estadística (CC)",
                       carrera="Ciencia de la Computación",
                       horas_conf=32, horas_cp=32, grupos_cp=2),
            Asignatura(id="EST-MAT", nombre="Estadística (Mat)",
                       carrera="Matemática", horas_conf=32, horas_cp=32,
                       grupos_cp=1),
        ),
    )


def _hoja():
    # 5 filas de carga + 2 de reserva -> filas 4..10 de la tabla.
    # 2 profesores + 1 de reserva -> filas 4..6 del panel.
    wb = Workbook()
    wb.remove(wb.active)
    depto = _departamento()
    construir_hoja_datos(wb, depto)
    construir_hoja_asignacion(wb, depto)
    return wb["Asignación"]


def _formulas_condicionales(ws):
    return [regla.formula[0]
            for reglas in ws.conditional_formatting._cf_rules.values()
            for regla in reglas if regla.formula]


def test_encabezados_del_panel():
    ws = _hoja()
    assert [ws[f"{c}3"].value for c in "JKLMN"] == [
        "Id", "Nombre", "Horas", "Asignaturas", "Tope"]


def test_el_panel_sale_del_claustro_y_no_de_una_copia():
    # Un profesor anadido a mano en la hoja Profesores aparece aqui solo. Si el
    # panel llevara los nombres escritos al generar, no aparecerian nunca.
    ws = _hoja()
    assert ws["J4"].value == '=IF(Profesores!A2="","",Profesores!A2)'
    assert ws["K4"].value == '=IF(Profesores!A2="","",Profesores!B2)'
    assert ws["N4"].value == '=IF(Profesores!A2="","",Profesores!D2)'
    # La segunda fila del panel mira la segunda del claustro.
    assert ws["J5"].value == '=IF(Profesores!A3="","",Profesores!A3)'


def test_total_de_horas():
    # Filtra por la columna de profesor (G) y suma la de horas (F), reserva
    # incluida: las filas creadas a mano cuentan igual.
    ws = _hoja()
    assert ws["L4"].value == '=IF($J4="","",SUMIF($G$4:$G$10,$J4,$F$4:$F$10))'


def test_asignaturas_distintas():
    # La Conf y los dos grupos de CP de la misma asignatura son 3 filas de
    # carga pero 1 asignatura. La cuenta se apoya en la columna de la hoja
    # Auxiliar que marca la primera aparicion de cada par profesor-asignatura.
    ws = _hoja()
    assert ws["M4"].value == (
        '=IF($J4="","",SUMIF($G$4:$G$10,$J4,AsignaturaNuevaProfesor))')


def test_una_fila_por_hueco_del_claustro():
    ws = _hoja()
    assert ws["J6"].value is not None    # la fila de reserva existe
    assert ws["J7"].value is None        # y no hay ninguna de mas


def test_alerta_de_sobrecarga_en_el_panel():
    # El mismo rojo que la fila TOTAL de `Carga por profesor`. La guarda del
    # tope vacio evita que dispare con la casilla en blanco, donde no hay con
    # que comparar.
    formulas = _formulas_condicionales(_hoja())
    assert any('$N4<>""' in f and "$L4>$N4" in f for f in formulas), formulas


def test_el_panel_queda_fuera_del_autofiltro():
    # Si entrara en el autofiltro, filtrar la tabla escondería filas del panel.
    assert _hoja().auto_filter.ref == "A3:H10"


def test_el_panel_esta_bloqueado():
    # Todo son formulas: se mira, no se escribe. Los datos se cambian en el
    # claustro y las horas en la propia tabla.
    ws = _hoja()
    for col in "JKLMN":
        assert ws[f"{col}4"].protection.locked is True, col


def test_la_canaleta_queda_vacia():
    ws = _hoja()
    assert ws["I3"].value is None
    assert ws["I4"].value is None


def test_las_cifras_del_panel_son_enteras():
    # Sin formato explicito, Calc mostraria "32,00" o "32" segun la
    # configuracion regional de quien abra el libro.
    ws = _hoja()
    for col in ("L", "M", "N"):
        assert ws[f"{col}4"].number_format == "0", col


def test_la_leyenda_explica_el_rojo_del_panel():
    ws = _hoja()
    textos = [ws[f"B{r}"].value or "" for r in range(11, 17)]
    assert any("tope" in t.lower() for t in textos), textos
