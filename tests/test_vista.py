from openpyxl import Workbook
from comun import vista


def test_ocultar_cuadricula():
    ws = Workbook().active
    vista.ocultar_cuadricula(ws)
    assert ws.sheet_view.showGridLines is False


def test_color_de_pestana():
    ws = Workbook().active
    vista.colorear_pestana(ws, "BBDEFB")
    assert ws.sheet_properties.tabColor.rgb.endswith("BBDEFB")


def test_zoom():
    ws = Workbook().active
    vista.fijar_zoom(ws, 110)
    assert ws.sheet_view.zoomScale == 110
