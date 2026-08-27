"""Tests del constructor generico de hoja de listado.

Las diez hojas de listado de los tres generadores salen de aqui, asi que lo que
se comprueba es el formato de la casa: donde van los encabezados, que haya una
fila por dato, y que la hoja quede protegida pero ordenable.
"""
from openpyxl import Workbook

from comun.hoja_listado import construir_hoja_listado, FILA_ENCABEZADO

ENCABEZADOS = ("Id", "Nombre", "Grado")
FILAS = [("PIAD", "Pedro I. Alonso Diaz", "Dr."),
         ("MARA", "Maria Ramirez", "MSc.")]


def _hoja(encabezados=ENCABEZADOS, filas=FILAS):
    wb = Workbook()
    wb.remove(wb.active)
    return construir_hoja_listado(wb, "Profesores", encabezados, filas,
                                  color_encabezado="D9D9D9")


def test_la_hoja_se_crea_con_el_nombre_pedido():
    ws = _hoja()
    assert ws.title == "Profesores"
    assert ws.parent.sheetnames == ["Profesores"]


def test_los_encabezados_van_en_la_primera_fila():
    ws = _hoja()
    assert [c.value for c in ws[FILA_ENCABEZADO][:3]] == list(ENCABEZADOS)


def test_hay_una_fila_por_dato_justo_debajo():
    ws = _hoja()
    assert ws["A2"].value == "PIAD"
    assert ws["C2"].value == "Dr."
    assert ws["A3"].value == "MARA"
    # Y nada mas: la fila siguiente esta vacia.
    assert ws["A4"].value is None


def test_los_encabezados_van_en_negrita_y_con_relleno():
    ws = _hoja()
    assert ws["A1"].font.bold is True
    assert ws["A1"].fill.start_color.rgb.endswith("D9D9D9")


def test_las_celdas_llevan_borde():
    ws = _hoja()
    # Perimetro medio, enrejado interno fino: la esquina superior izquierda
    # tiene el lado externo arriba y a la izquierda.
    assert ws["A1"].border.top.style == "medium"
    assert ws["A1"].border.left.style == "medium"
    assert ws["B2"].border.top.style == "thin"


def test_la_fila_de_encabezado_queda_congelada():
    ws = _hoja()
    assert ws.freeze_panes == "A2"


def test_la_hoja_va_protegida_pero_se_puede_ordenar():
    ws = _hoja()
    assert ws.protection.sheet is True
    # En SheetProtection cada atributo significa "esta operacion queda
    # prohibida": ordenar habilitado es sort=False.
    assert ws.protection.sort is False
    assert ws.protection.autoFilter is False


def test_la_hoja_no_muestra_la_cuadricula_de_fondo():
    ws = _hoja()
    assert ws.sheet_view.showGridLines is False


def test_las_columnas_se_autoajustan_al_texto_mas_largo():
    ws = _hoja()
    # "Pedro I. Alonso Diaz" son 20 caracteres; la columna tiene que caberlo.
    assert ws.column_dimensions["B"].width >= 20


def test_las_filas_de_datos_llevan_padding():
    ws = _hoja()
    assert ws["A2"].alignment.vertical == "center"
    assert ws["A2"].alignment.indent == 1
    assert ws.row_dimensions[2].height is not None


def test_un_listado_sin_filas_deja_solo_los_encabezados():
    # Un generador puede quedarse sin datos de una entidad (p. ej. una facultad
    # sin locales declarados); la hoja tiene que salir igual y no reventar.
    ws = _hoja(filas=[])
    assert [c.value for c in ws[1][:3]] == list(ENCABEZADOS)
    assert ws["A2"].value is None
