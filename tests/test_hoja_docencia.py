"""Tests de las hojas Profesores y Docencia del generador de horarios.

Las dos existen solo si el YAML declara profesores. Un YAML sin la seccion
-todos los anteriores a la fase 2- no debe ganar dos pestanas vacias.
"""
from openpyxl import Workbook

from horarios.hoja_profesores import construir_hoja_profesores
from horarios.hoja_docencia import construir_hoja_docencia
from horarios.modelo import (
    Asignatura, Anio, Grupo, Profesor, Docencia, Facultad,
)

ANIOS = {"C1": Anio(carrera="C", numero=1, asignaturas=(
    Asignatura(id="AMI-C", nombre="Análisis Matemático I (Conf)", frecuencia=1),
    Asignatura(id="AMI-CP", nombre="Análisis Matemático I (C.P.)", frecuencia=2)))}

FACULTAD = Facultad(
    aulas=("Aula 1",), dias=("Lunes",), turnos=6,
    grupos=(Grupo(carrera="C", anio=1, sesion=1, numero=1),
            Grupo(carrera="C", anio=1, sesion=1, numero=2)),
    anios=ANIOS,
    profesores=(Profesor(id="PIAD", nombre="Pedro I. Alonso Diaz", grado="Dr.",
                         tope_turnos=160),
                Profesor(id="MARA", nombre="Maria Ramirez", grado="MSc.")),
    docencia=(Docencia(grupo="C111", asignatura="AMI-C", profesor="PIAD"),
              Docencia(grupo="C112", asignatura="AMI-CP", profesor="MARA")),
)


def _hoja(constructor, facultad=FACULTAD):
    wb = Workbook()
    wb.remove(wb.active)
    constructor(wb, facultad)
    return wb[wb.sheetnames[0]] if wb.sheetnames else None


def test_profesores_lista_el_claustro_con_grado_y_tope():
    ws = _hoja(construir_hoja_profesores)
    assert ws.title == "Profesores"
    assert [c.value for c in ws[1][:4]] == ["Id", "Nombre", "Grado", "Tope de turnos"]
    assert ws["A2"].value == "PIAD"
    assert ws["D2"].value == 160
    # Sin tope declarado la casilla queda en blanco, no en cero: un 0 se leeria
    # como "no puede impartir nada".
    assert ws["D3"].value == ""


def test_docencia_dice_quien_imparte_que_en_cada_grupo():
    ws = _hoja(construir_hoja_docencia)
    assert ws.title == "Docencia"
    assert [c.value for c in ws[1][:4]] == ["Grupo", "Asignatura", "Profesor",
                                            "Nombre"]
    assert ws["A2"].value == "C111"
    assert ws["B2"].value == "AMI-C"
    assert ws["C2"].value == "PIAD"
    # El nombre completo al lado del id: el resto del libro trabaja con ids y
    # esta es la hoja donde se comprueba a quien corresponden. Lo resuelve una
    # formula, para que siga al profesor que se escriba (ver el test de abajo).
    assert "VLOOKUP" in ws["D2"].value


def test_sin_profesores_no_se_crea_ninguna_de_las_dos_hojas():
    vacia = Facultad(aulas=("Aula 1",), dias=("Lunes",), turnos=6,
                     grupos=FACULTAD.grupos, anios=ANIOS)
    for constructor in (construir_hoja_profesores, construir_hoja_docencia):
        wb = Workbook()
        wb.remove(wb.active)
        constructor(wb, vacia)
        assert wb.sheetnames == []


def test_las_dos_hojas_quedan_protegidas():
    for constructor in (construir_hoja_profesores, construir_hoja_docencia):
        ws = _hoja(constructor)
        assert ws.protection.sheet is True, ws.title


def test_el_claustro_cuenta_los_turnos_semanales_de_cada_profesor():
    ws = _hoja(construir_hoja_profesores)
    assert [c.value for c in ws[1][:5]] == ["Id", "Nombre", "Grado",
                                            "Tope de turnos", "Turnos semanales"]
    # Va como formula: si se cambia quien imparte algo en la hoja Docencia, el
    # total tiene que seguirlo sin regenerar el libro.
    assert str(ws["E2"].value).startswith("=")
    assert "SUMIF" in ws["E2"].value


def test_hay_alerta_cuando_los_turnos_pasan_del_tope():
    ws = _hoja(construir_hoja_profesores)
    formulas = [r.formula[0] for reglas in ws.conditional_formatting._cf_rules.values()
                for r in reglas if r.formula]
    assert any("$D" in f and "$E" in f and ">" in f for f in formulas), formulas


def test_el_nombre_de_la_docencia_sigue_al_profesor_que_se_escriba():
    """Hallazgo de la revision: 'Nombre' era una copia escrita al generar, y
    bloqueada, al lado de una columna 'Profesor' editable. Cambiar el profesor
    dejaba el nombre del anterior, sin forma de corregirlo."""
    ws = _hoja(construir_hoja_docencia)
    assert str(ws["D2"].value).startswith("=")
    assert "ProfesoresTabla" in ws["D2"].value


def test_las_columnas_editables_de_docencia_tienen_desplegable():
    """Hallazgo de la revision: se escribia a mano sin lista ni validacion, y un
    id mal escrito dejaba la celda de profesor de la rejilla en blanco sin decir
    por que."""
    ws = _hoja(construir_hoja_docencia)
    fuentes = {dv.formula1 for dv in ws.data_validations.dataValidation}
    assert {"GruposValidos", "AsigId", "ProfesoresValidos"} <= fuentes, fuentes


def test_un_profesor_fuera_de_la_lista_se_resalta_en_docencia():
    ws = _hoja(construir_hoja_docencia)
    formulas = [r.formula[0] for reglas in ws.conditional_formatting._cf_rules.values()
                for r in reglas if r.formula]
    assert any("ProfesoresValidos" in f and "COUNTIF" in f for f in formulas), formulas


def test_las_filas_de_reserva_ya_traen_sus_formulas():
    """Se vio al mirar el PDF: al escribir una fila nueva en la reserva, la
    columna calculada se quedaba en blanco. Las lineas libres estaban preparadas
    a medias: con borde y alto, pero sin la formula que las hace utiles."""
    ws = _hoja(construir_hoja_docencia)
    fila_libre = next(f for f in range(2, ws.max_row + 1) if ws[f"A{f}"].value is None)
    assert str(ws[f"D{fila_libre}"].value).startswith("="), fila_libre


def test_el_claustro_de_horarios_calcula_tambien_en_la_reserva():
    ws = _hoja(construir_hoja_profesores)
    fila_libre = next(f for f in range(2, ws.max_row + 1) if ws[f"A{f}"].value is None)
    assert str(ws[f"E{fila_libre}"].value).startswith("="), fila_libre
    # Y con la fila vacia no muestra un 0 suelto.
    assert 'IF(' in ws[f"E{fila_libre}"].value
