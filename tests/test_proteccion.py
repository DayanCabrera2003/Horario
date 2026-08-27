"""Tests del modulo de proteccion de hojas."""
from openpyxl import Workbook
from openpyxl.styles import Font

from comun import proteccion


def _hoja():
    wb = Workbook()
    return wb.active


def test_proteger_hoja_marca_la_hoja_como_protegida():
    ws = _hoja()
    proteccion.proteger_hoja(ws)
    assert ws.protection.sheet is True


def test_los_rangos_declarados_quedan_editables():
    ws = _hoja()
    proteccion.proteger_hoja(ws, editables=["B2:C3"])
    assert ws["B2"].protection.locked is False
    assert ws["C3"].protection.locked is False


def test_lo_no_declarado_sigue_bloqueado():
    ws = _hoja()
    proteccion.proteger_hoja(ws, editables=["B2:C3"])
    assert ws["A1"].protection.locked is True


def test_acepta_rangos_multiples_separados_por_espacios():
    # horarios/layout.py devuelve "B4:F4 B6:F6", que ws[...] rechaza.
    ws = _hoja()
    proteccion.proteger_hoja(ws, editables=["B4:C4 B6:C6"])
    assert ws["B4"].protection.locked is False
    assert ws["B6"].protection.locked is False


def test_desbloquear_no_pisa_los_demas_estilos():
    ws = _hoja()
    ws["B2"].font = Font(bold=True)
    proteccion.proteger_hoja(ws, editables=["B2:B2"])
    assert ws["B2"].font.bold is True
    assert ws["B2"].protection.locked is False


def test_desbloquear_no_protege_la_hoja_por_si_solo():
    # `desbloquear` es parte de la interfaz publica del modulo y se usa por
    # separado de `proteger_hoja`; su contrato se fija aqui.
    ws = _hoja()
    proteccion.desbloquear(ws, ["B2:C3"])
    assert ws["B2"].protection.locked is False
    assert ws.protection.sheet is False


def test_permitir_orden_deja_ordenar():
    # En OOXML cada atributo significa "esta operacion queda prohibida":
    # ponerlos en False es lo que las habilita.
    ws = _hoja()
    proteccion.proteger_hoja(ws, permitir_orden=True)
    assert ws.protection.sort is False


def test_permitir_filtro_deja_autofiltrar():
    ws = _hoja()
    proteccion.proteger_hoja(ws, permitir_filtro=True)
    assert ws.protection.autoFilter is False


def test_ordenar_y_filtrar_se_conceden_por_separado():
    # Filtrar solo esconde filas; ordenar las mueve de sitio. En una hoja cuyas
    # filas tienen gemela por posicion en otra hoja, filtrar es inocuo pero
    # ordenar descuadra la gemela en silencio, asi que no van juntas.
    ws = _hoja()
    proteccion.proteger_hoja(ws, permitir_filtro=True)
    assert ws.protection.sort is True

    otra = _hoja()
    proteccion.proteger_hoja(otra, permitir_orden=True)
    assert otra.protection.autoFilter is True


def test_sin_permisos_ordenar_y_filtrar_quedan_prohibidas():
    ws = _hoja()
    proteccion.proteger_hoja(ws)
    assert ws.protection.sort is True
    assert ws.protection.autoFilter is True


def test_celda_unica_duplica_la_referencia():
    # ws["F4"] devuelve una Cell, no una tupla de filas; el rango duplicado
    # es lo que espera `desbloquear`.
    assert proteccion.celda_unica("F4") == "F4:F4"
