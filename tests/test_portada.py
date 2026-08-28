import datetime

from openpyxl import Workbook

from comun import portada


def _wb():
    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Asignación")
    wb.create_sheet("Profesores")
    wb.create_sheet("Aulas")
    return wb


HOJAS = (
    ("Asignación", "Quién imparte cada conferencia", portada.SE_ESCRIBE),
    ("Profesores", "Cuántas horas acumula cada uno", portada.SE_CALCULA),
    ("Aulas", "Las aulas de la facultad", portada.SON_DATOS),
)

GENERADO = datetime.datetime(2026, 8, 25, 14, 30)


def _construir(wb, **kwargs):
    """Llama a construir_portada con los datos de ejemplo, dejando sobrescribir
    solo lo que cada test necesita variar."""
    datos = dict(titulo="Gestión", subtitulo="Mat · 2026",
                 origen="config/departamento.yaml", hojas=HOJAS,
                 generado=GENERADO)
    datos.update(kwargs)
    portada.construir_portada(wb, **datos)
    return wb[portada.NOMBRE_HOJA]


def _textos(ws):
    """Todos los textos de la hoja, para buscar sin depender de coordenadas."""
    return [c.value for fila in ws.iter_rows() for c in fila
            if isinstance(c.value, str)]


def _fila_del_indice(ws, nombre: str):
    """La fila del indice que lista a la hoja `nombre`, o None si no esta."""
    for fila in ws.iter_rows():
        if fila[0].value == nombre:
            return fila
    return None


def test_la_portada_es_la_primera_hoja():
    wb = _wb()
    _construir(wb)
    assert wb.sheetnames[0] == "Portada"


def test_muestra_el_archivo_de_origen():
    ws = _construir(_wb())
    assert any("config/departamento.yaml" in t for t in _textos(ws))


def test_la_fecha_es_inyectable_para_que_la_salida_sea_reproducible():
    # Sin inyectarla, dos corridas producen libros distintos y cualquier test
    # que compare archivos empieza a fallar por el reloj. Se compara la fecha
    # completa y no solo el ano: si saliera del reloj, el dia no coincidiria.
    ws = _construir(_wb())
    assert any("25/08/2026 14:30" in t for t in _textos(ws))


def test_el_indice_enlaza_a_cada_hoja():
    ws = _construir(_wb())
    enlaces = [c.hyperlink.location for fila in ws.iter_rows() for c in fila
               if c.hyperlink is not None]
    # `location` y no `target`: un enlace interno va en el atributo location del
    # XML. Asignar la cadena directamente lo guardaria como relacion externa.
    assert any("Asignación" in e for e in enlaces)
    assert any("Profesores" in e for e in enlaces)


def test_el_indice_marca_que_hojas_se_escriben_a_mano():
    # Se mira la columna del indice, no el texto suelto de la hoja: la frase de
    # las instrucciones ya contiene "se escriben" y "se calculan", asi que
    # buscarlas en toda la hoja pasaria aunque el indice no marcara nada.
    ws = _construir(_wb())
    assert _fila_del_indice(ws, "Asignación")[2].value == "se escribe"
    assert _fila_del_indice(ws, "Profesores")[2].value == "se calcula"


def test_el_indice_distingue_las_hojas_de_datos():
    # Un listado no es ninguna de las dos cosas: ni se calcula solo ni es donde
    # se planifica. Es la respuesta a "¿dónde están los datos?".
    ws = _construir(_wb())
    assert _fila_del_indice(ws, "Aulas")[2].value == "datos del problema"


def test_el_indice_describe_cada_hoja():
    ws = _construir(_wb())
    assert _fila_del_indice(ws, "Asignación")[1].value == "Quién imparte cada conferencia"


def test_la_portada_queda_protegida():
    ws = _construir(_wb())
    assert ws.protection.sheet is True


def test_la_portada_avisa_de_no_dejar_huecos_en_las_listas():
    # Comprobado en Calc: COUNTA cuenta celdas no vacias, asi que un hueco en
    # medio de un listado deja fuera del rango a lo que queda debajo, y ese
    # dato desaparece del desplegable sin avisar de nada.
    ws = _construir(_wb())
    texto = " ".join(_textos(ws))
    assert "en blanco" in texto
    assert "desplegable" in texto


def test_los_parrafos_de_instrucciones_caben_en_la_pagina():
    """Se veian cortados al imprimir la portada.

    Son dos parrafos de 180 y 220 caracteres en una columna de 45: sin combinar
    ni ajustar el texto, lo que no cabe se pierde en el papel y empuja la hoja a
    una segunda pagina con una sola columna suelta.
    """
    ws = _construir(_wb())
    filas = [f for f in range(1, 15)
             if isinstance(ws[f"A{f}"].value, str) and len(ws[f"A{f}"].value) > 100]
    assert filas, "no se encontraron los parrafos"
    combinadas = {str(r) for r in ws.merged_cells.ranges}
    for fila in filas:
        assert ws[f"A{fila}"].alignment.wrap_text is True, fila
        assert any(str(fila) in c for c in combinadas), (fila, combinadas)
        assert ws.row_dimensions[fila].height, fila


def test_un_origen_largo_no_ensancha_la_columna_de_etiquetas():
    # La ruta del YAML puede ser larguisima; la columna A es la de etiquetas y
    # nombres de hoja, y no tiene por que crecer con ella.
    ws = _construir(_wb(), origen="/una/ruta/absurdamente/larga/" + "x" * 80)
    assert ws.column_dimensions["A"].width < 30


def test_unas_instrucciones_vacias_no_dejan_una_fila_combinada_de_mas():
    # `instrucciones` es un parametro: un generador puede no querer ninguna.
    ws = _construir(_wb(), instrucciones="")
    combinadas = {str(r) for r in ws.merged_cells.ranges}
    assert not any(str(portada._FILA_INSTRUCCIONES + 1) in c for c in combinadas)
