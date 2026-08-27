from openpyxl import Workbook
from comun import impresion


def test_apaisado():
    ws = Workbook().active
    impresion.preparar(ws)
    assert ws.page_setup.orientation == "landscape"


def test_ajuste_a_una_pagina_de_ancho():
    ws = Workbook().active
    impresion.preparar(ws)
    # fitToWidth no hace nada si fitToPage no esta activo: esa es la trampa.
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True
    assert ws.page_setup.fitToWidth == 1
    assert ws.page_setup.fitToHeight == 0


def test_repetir_encabezado_en_cada_pagina():
    ws = Workbook().active
    impresion.preparar(ws, filas_encabezado="1:3")
    # openpyxl normaliza el rango a referencias absolutas al leerlo de vuelta
    # ("1:3" se guarda pero se devuelve como "$1:$3"); no es cosa nuestra.
    assert ws.print_title_rows == "$1:$3"


def test_no_pisa_otros_ajustes_de_pagina_ya_fijados():
    # El modulo muta el objeto `pageSetUpPr` en vez de reemplazarlo justamente
    # para esto: si otro modulo ya fijo `autoPageBreaks`, tiene que sobrevivir.
    ws = Workbook().active
    ws.sheet_properties.pageSetUpPr.autoPageBreaks = False
    impresion.preparar(ws)
    assert ws.sheet_properties.pageSetUpPr.autoPageBreaks is False
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True


def test_funciona_aunque_la_hoja_no_traiga_ajustes_de_pagina():
    # openpyxl siempre entrega un `pageSetUpPr`, tanto en una hoja nueva como en
    # una recien cargada de un archivo, asi que este caso no se da hoy. La guarda
    # existe por si eso cambia, y aqui se comprueba que hace lo suyo.
    ws = Workbook().active
    ws.sheet_properties.pageSetUpPr = None
    impresion.preparar(ws)
    assert ws.sheet_properties.pageSetUpPr is not None
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True
