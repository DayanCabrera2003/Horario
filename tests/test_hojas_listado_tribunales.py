"""Tests de las hojas de listado del generador de tribunales.

Son los datos del problema a la vista: quienes son los profesores, quienes los
estudiantes, que locales hay y que dias con que momentos. Lo que mas importa es
que aparezcan **todos**, participen o no en una tesis: hoy un profesor sin
tribunal o un estudiante sin tesis no salen por ninguna parte del libro.
"""
from openpyxl import Workbook

from tribunales.hoja_profesores import construir_hoja_profesores
from tribunales.hoja_estudiantes import construir_hoja_estudiantes
from tribunales.hoja_locales import construir_hoja_locales
from tribunales.hoja_dias import construir_hoja_dias
from tribunales.modelo import (
    Profesor, Estudiante, Local, Momento, Dia, Tesis, Facultad,
)

# Una facultad con un profesor y un estudiante que no participan en ninguna
# tesis: son el caso que estas hojas existen para cubrir.
FACULTAD = Facultad(
    profesores=(Profesor(id="PIAD", nombre="Pedro I. Alonso Diaz", grado="Dr."),
                Profesor(id="MARA", nombre="Maria Ramirez", grado="MSc."),
                Profesor(id="LGOM", nombre="Luis Gomez", grado="Dr."),
                Profesor(id="ANSU", nombre="Ana Suarez", grado=""),
                Profesor(id="SOLO", nombre="Sin Tribunal", grado="MSc.")),
    estudiantes=(Estudiante(id="JPER", nombre="Juan Perez"),
                 Estudiante(id="NADI", nombre="Nadie Sin Tesis")),
    locales=(Local(id="POST", nombre="Postgrado"),
             Local(id="DECA", nombre="Decanato")),
    dias=(Dia(fecha="2026-07-27", momentos=(Momento("09:00", "10:00"),
                                            Momento("10:00", "11:00"))),
          Dia(fecha="2026-07-28", momentos=(Momento("09:00", "10:00"),))),
    tesis=(Tesis(estudiantes=("JPER",), tutores=("PIAD",), oponente="MARA",
                 presidente="LGOM", secretario="ANSU"),),
)


def _hoja(constructor):
    wb = Workbook()
    wb.remove(wb.active)
    constructor(wb, FACULTAD)
    return wb[wb.sheetnames[0]]


def _columna(ws, letra, desde=2):
    return [ws[f"{letra}{f}"].value
            for f in range(desde, ws.max_row + 1)
            if ws[f"{letra}{f}"].value is not None]


def test_profesores_lista_id_nombre_y_grado():
    ws = _hoja(construir_hoja_profesores)
    assert ws.title == "Profesores"
    assert [c.value for c in ws[1][:3]] == ["Id", "Nombre", "Grado"]
    assert ws["A2"].value == "PIAD"
    assert ws["B2"].value == "Pedro I. Alonso Diaz"
    assert ws["C2"].value == "Dr."


def test_profesores_incluye_a_quien_no_esta_en_ningun_tribunal():
    ws = _hoja(construir_hoja_profesores)
    assert "SOLO" in _columna(ws, "A")
    assert len(_columna(ws, "A")) == 5


def test_profesores_no_lleva_tope_de_horas():
    # Divergencia declarada en el diseno: un tribunal de tesis no es carga
    # docente, asi que el tope vive solo en el libro del departamento.
    ws = _hoja(construir_hoja_profesores)
    assert all("ope" not in str(c.value) for c in ws[1] if c.value)


def test_estudiantes_incluye_a_quien_no_tiene_tesis():
    ws = _hoja(construir_hoja_estudiantes)
    assert ws.title == "Estudiantes"
    assert [c.value for c in ws[1][:2]] == ["Id", "Nombre"]
    assert _columna(ws, "A") == ["JPER", "NADI"]


def test_locales_lista_id_y_nombre():
    ws = _hoja(construir_hoja_locales)
    assert ws.title == "Locales"
    assert [c.value for c in ws[1][:2]] == ["Id", "Nombre"]
    assert _columna(ws, "B") == ["Postgrado", "Decanato"]


def test_dias_muestra_la_fecha_su_hoja_y_sus_momentos():
    ws = _hoja(construir_hoja_dias)
    assert ws.title == "Días"
    assert [c.value for c in ws[1][:3]] == ["Fecha", "Hoja", "Momentos"]
    assert ws["A2"].value == "2026-07-27"
    # La fecha ISO es la que se escribe en el YAML; el nombre de hoja es otro.
    # Que la equivalencia este escrita evita tener que deducirla.
    assert ws["B2"].value == "27 jul (lun)"
    assert ws["C2"].value == "09:00-10:00, 10:00-11:00"


def test_los_listados_quedan_protegidos_y_ordenables():
    for constructor in (construir_hoja_profesores, construir_hoja_estudiantes,
                        construir_hoja_locales, construir_hoja_dias):
        ws = _hoja(constructor)
        assert ws.protection.sheet is True, ws.title
        assert ws.protection.sort is False, ws.title
