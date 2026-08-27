"""Los listados que son fuente de algo se editan; los demas no.

Un listado desbloqueado que no alimenta ninguna formula seria mentira: se podria
escribir en el y no pasaria nada. Solo se desbloquean los tres que si mandan:
`Aulas` (desplegable de aula), `Profesores` del departamento (desplegable de
profesor y tope) y `Estudiantes` (desplegable de estudiante).
"""
from openpyxl import Workbook

from departamento.hoja_profesores import construir_hoja_profesores
from departamento.hoja_carga import construir_hoja_carga
from departamento.hoja_asignaturas import construir_hoja_asignaturas
from departamento.modelo import Profesor, Asignatura, Departamento
from horarios.hoja_aulas import construir_hoja_aulas
from horarios.hoja_estructura import construir_hoja_estructura
from horarios.modelo import Asignatura as AsigH, Anio, Grupo, Facultad

DEPTO = Departamento(
    nombre="D", semestre="1", tope_horas=160, filas_por_profesor=4,
    profesores=(Profesor(id="PIAD", nombre="Pedro", grado="Dr."),
                Profesor(id="ANSU", nombre="Ana", grado="MSc.", tope_horas=80)),
    asignaturas=(Asignatura(id="EST", nombre="Estadística", carrera="CC",
                            horas_conf=32, horas_cp=32, grupos_cp=1),),
)

FACULTAD = Facultad(
    aulas=("Aula 1", "Lab"), dias=("Lunes",), turnos=6,
    grupos=(Grupo(carrera="C", anio=1, sesion=1, numero=1),),
    anios={"C1": Anio(carrera="C", numero=1, asignaturas=(
        AsigH(id="L-C", nombre="Lógica", frecuencia=1),))},
)


def _hoja(constructor, datos):
    wb = Workbook()
    wb.remove(wb.active)
    constructor(wb, datos)
    return wb[wb.sheetnames[0]]


def test_el_listado_de_aulas_se_puede_editar():
    ws = _hoja(construir_hoja_aulas, FACULTAD)
    assert ws.protection.sheet is True
    # La primera fila libre de la reserva tambien: es donde se anade.
    assert ws["A2"].protection.locked is False
    assert ws["A4"].protection.locked is False


def test_un_listado_que_no_alimenta_nada_sigue_bloqueado():
    # Escribir aqui no cambiaria ninguna formula del libro.
    ws = _hoja(construir_hoja_estructura, FACULTAD)
    assert ws["B2"].protection.locked is not False


def test_el_claustro_se_edita_entero_menos_lo_calculado():
    ws = _hoja(construir_hoja_profesores, DEPTO)
    # Anadir un profesor es escribir id, nombre, grado y tope en una fila libre.
    for col in ("A", "B", "C", "D"):
        assert ws[f"{col}2"].protection.locked is False, col
        assert ws[f"{col}5"].protection.locked is False, col
    # El origen del tope no: lo calcula una formula.
    assert ws["E2"].protection.locked is not False


def test_el_origen_del_tope_se_recalcula_al_editarlo():
    # Es formula y no texto escrito: si se edita el tope, la etiqueta tiene que
    # seguirlo en vez de quedarse mintiendo.
    ws = _hoja(construir_hoja_profesores, DEPTO)
    assert str(ws["E2"].value).startswith("=")
    assert "160" in ws["E2"].value


def test_el_tope_del_reporte_de_carga_vuelve_a_estar_bloqueado():
    # Se editaba aqui como arruga declarada de la fase 1. Su sitio es el
    # listado, y aqui pasa a ser una copia calculada.
    ws = _hoja(construir_hoja_carga, DEPTO)
    assert ws["D4"].protection.locked is not False
    assert str(ws["D4"].value).startswith("=")
    assert "ProfesoresTabla" in ws["D4"].value


def test_el_reporte_de_carga_lee_el_nombre_del_claustro():
    ws = _hoja(construir_hoja_carga, DEPTO)
    assert "ProfesoresTabla" in str(ws["B4"].value)


def test_la_alerta_de_sobrecarga_sigue_comparando_total_contra_tope():
    ws = _hoja(construir_hoja_carga, DEPTO)
    formulas = [r.formula[0] for reglas in ws.conditional_formatting._cf_rules.values()
                for r in reglas if r.formula]
    assert any("$D$" in f and ">" in f for f in formulas)


def test_el_plan_de_asignaturas_no_deja_tocar_el_total():
    ws = _hoja(construir_hoja_asignaturas, DEPTO)
    assert ws["G2"].protection.locked is not False
